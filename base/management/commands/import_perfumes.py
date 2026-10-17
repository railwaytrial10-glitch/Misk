import os
import re
from decimal import Decimal
from django.core.management.base import BaseCommand
from django.conf import settings
from base.models import Product, ProductVariant, ProductImage, Category
import openpyxl
import cloudinary.uploader


class Command(BaseCommand):
    help = "Import perfumes from Excel sheet and upload local images to Cloudinary"

    def add_arguments(self, parser):
        parser.add_argument(
            "--file",
            type=str,
            default="Misk_Perfumes_Clean_ALLDONE.xlsx",
            help="Path to the Excel file",
        )
        parser.add_argument(
            "--images-dir",
            type=str,
            default="images",
            help="Directory containing the perfume images",
        )
        parser.add_argument(
            "--skip-images",
            action="store_true",
            help="Skip uploading images to Cloudinary",
        )
        parser.add_argument(
            "--limit",
            type=int,
            default=0,
            help="Limit number of rows to import (0 for all)",
        )

    def handle(self, *args, **options):
        file_path = options["file"]
        images_dir = options["images_dir"]
        skip_images = options["skip_images"]
        limit = options["limit"]

        if not os.path.exists(file_path):
            self.stderr.write(self.style.ERROR(f"Excel file not found at: {file_path}"))
            return

        self.stdout.write(self.style.NOTICE(f"Loading Excel file: {file_path}..."))
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active

        # Map images by row number suffix, e.g. 'sauvage-2.jpg' -> 2
        image_map = {}
        if os.path.exists(images_dir) and not skip_images:
            for fname in os.listdir(images_dir):
                match = re.search(r"-(\d+)\.(jpg|jpeg|png|webp)$", fname, re.IGNORECASE)
                if match:
                    row_num = int(match.group(1))
                    image_map[row_num] = os.path.join(images_dir, fname)
            self.stdout.write(
                self.style.SUCCESS(f"Found {len(image_map)} images in '{images_dir}' mapped to rows.")
            )
        elif not skip_images:
            self.stdout.write(
                self.style.WARNING(f"Images folder '{images_dir}' not found. Continuing without images.")
            )

        total_rows = ws.max_row
        self.stdout.write(f"Total rows to process: {total_rows - 1}")

        success_count = 0
        error_count = 0
        processed_count = 0

        for row_idx in range(2, total_rows + 1):
            if limit > 0 and processed_count >= limit:
                break

            processed_count += 1

            # Extract cell values
            name_en = ws.cell(row=row_idx, column=1).value
            name_ar = ws.cell(row=row_idx, column=2).value
            desc_en = ws.cell(row=row_idx, column=3).value
            desc_ar = ws.cell(row=row_idx, column=4).value

            if not name_en:
                continue

            name_en = str(name_en).strip()
            name_ar = str(name_ar).strip() if name_ar else name_en
            desc_en = str(desc_en).strip() if desc_en else ""
            desc_ar = str(desc_ar).strip() if desc_ar else desc_en

            cat_en = ws.cell(row=row_idx, column=5).value
            cat_ar = ws.cell(row=row_idx, column=6).value
            conc_en = ws.cell(row=row_idx, column=7).value
            conc_ar = ws.cell(row=row_idx, column=8).value
            frag_fam_en = ws.cell(row=row_idx, column=9).value
            frag_fam_ar = ws.cell(row=row_idx, column=10).value
            is_active_val = ws.cell(row=row_idx, column=11).value
            is_bestseller_val = ws.cell(row=row_idx, column=12).value

            is_active = True if is_active_val in [True, "TRUE", "True", 1, "1", None] else False
            is_bestseller = True if is_bestseller_val in [True, "TRUE", "True", 1, "1"] else False

            # 2. Check if product already exists -> SKIP IT
            existing_product = Product.objects.filter(name__iexact=name_en).first()
            if existing_product:
                self.stdout.write(
                    self.style.WARNING(
                        f"[{row_idx}/{total_rows}] Skipped (Already exists in DB): {name_en}"
                    )
                )
                continue

            try:
                # 1. Resolve or Create Category
                category_obj = None
                target_cat_name = None
                if cat_ar:
                    target_cat_name = str(cat_ar).strip()
                elif cat_en:
                    target_cat_name = str(cat_en).strip()

                if target_cat_name:
                    # Match existing categories by English name or Arabic name
                    category_obj = (
                        Category.objects.filter(name__iexact=target_cat_name).first()
                        or Category.objects.filter(name_ar__iexact=target_cat_name).first()
                    )
                    if not category_obj:
                        category_obj = Category.objects.create(
                            name=str(cat_en).strip() if cat_en else target_cat_name,
                            name_ar=str(cat_ar).strip() if cat_ar else target_cat_name,
                        )

                # 2. Create New Product
                product = Product.objects.create(
                    name=name_en,
                    name_ar=name_ar,
                    description=desc_en,
                    description_ar=desc_ar,
                    concentration=str(conc_en).strip() if conc_en else "",
                    concentration_ar=str(conc_ar).strip() if conc_ar else "",
                    fragrance_family=str(frag_fam_en).strip() if frag_fam_en else "",
                    fragrance_family_ar=str(frag_fam_ar).strip() if frag_fam_ar else "",
                    is_active=is_active,
                    is_bestseller=is_bestseller,
                )

                if category_obj:
                    product.categories.add(category_obj)

                # 3. Process Variants (3 pairs of Volume & Price)
                # Variant 1: cols 13-16, Variant 2: cols 18-21, Variant 3: cols 23-26
                variants_configs = [
                    (13, 14, 15, 16),
                    (18, 19, 20, 21),
                    (23, 24, 25, 26),
                ]

                created_variants = []

                for vol_col, price_col, cmp_price_col, stock_col in variants_configs:
                    raw_vol = ws.cell(row=row_idx, column=vol_col).value
                    raw_price = ws.cell(row=row_idx, column=price_col).value

                    if raw_vol is None or raw_price is None:
                        continue

                    # Format volume (e.g., 30 -> '30ml')
                    vol_str = str(raw_vol).strip()
                    if not vol_str.lower().endswith("ml"):
                        vol_str = f"{vol_str}ml"

                    try:
                        price_dec = Decimal(str(raw_price).strip())
                    except Exception:
                        continue

                    # Compare at price
                    raw_cmp = ws.cell(row=row_idx, column=cmp_price_col).value
                    cmp_price = None
                    if raw_cmp:
                        try:
                            cmp_price = Decimal(str(raw_cmp).strip())
                        except Exception:
                            pass

                    # Stock
                    raw_stock = ws.cell(row=row_idx, column=stock_col).value
                    stock_qty = 0
                    if raw_stock is not None:
                        try:
                            stock_qty = int(raw_stock)
                        except Exception:
                            stock_qty = 100
                    else:
                        stock_qty = 100

                    variant, _ = ProductVariant.objects.update_or_create(
                        product=product,
                        volume=vol_str,
                        defaults={
                            "price": price_dec,
                            "compare_at_price": cmp_price,
                            "stock": stock_qty,
                            "is_active": True,
                        },
                    )
                    created_variants.append(variant)

                # 4. Upload Image to Cloudinary & Attach to all variants
                if not skip_images and row_idx in image_map and created_variants:
                    img_path = image_map[row_idx]
                    public_id = f"products/{product.id}_{os.path.splitext(os.path.basename(img_path))[0]}"

                    try:
                        upload_result = cloudinary.uploader.upload(
                            img_path,
                            public_id=public_id,
                            overwrite=True,
                            resource_type="image",
                        )
                        cloudinary_image_id = upload_result.get("public_id")

                        for variant in created_variants:
                            # Avoid duplicates: update or create thumbnail image
                            existing_image = variant.images.filter(is_thumbnail=True).first()
                            if existing_image:
                                existing_image.img = cloudinary_image_id
                                existing_image.save()
                            else:
                                ProductImage.objects.create(
                                    variant=variant,
                                    img=cloudinary_image_id,
                                    is_thumbnail=True,
                                )
                    except Exception as img_err:
                        self.stderr.write(
                            self.style.WARNING(
                                f"Failed to upload image for '{name_en}' (row {row_idx}): {img_err}"
                            )
                        )

                success_count += 1
                self.stdout.write(
                    self.style.SUCCESS(
                        f"[{row_idx}/{total_rows}] Imported: {name_en} ({len(created_variants)} variants)"
                    )
                )

            except Exception as e:
                error_count += 1
                self.stderr.write(self.style.ERROR(f"Error at row {row_idx} ({name_en}): {e}"))

        self.stdout.write(
            self.style.SUCCESS(
                f"\nFinished! Successfully processed {success_count} perfumes ({error_count} errors)."
            )
        )
